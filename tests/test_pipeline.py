import pytest
from openpyxl import load_workbook

from mlrrecon.db import migrate
from mlrrecon.ingest import ingest_zip, parse_column
from mlrrecon.recon import reconcile, set_status
from mlrrecon.report import peer_view, write_workbook

from .conftest import build_zip, default_data


def test_migrations_apply_once_in_order(con):
    assert migrate(con) == []
    versions = [r[0] for r in con.execute("SELECT version FROM schema_migrations ORDER BY version")]
    assert versions == ["001_init", "002_exceptions", "003_market_mlr_view", "004_exception_indexes"]


def test_parse_column():
    assert parse_column("cmm_small_group_cy") == ("cmm_small_group", "cy")
    assert parse_column("mini_med_large_group_total") == ("mini_med_large_group", "total")
    assert parse_column("created_date") is None


def test_ingest_is_idempotent(con, puf):
    first = ingest_zip(con, puf)
    assert first["year"] == 2024 and first["filings"] == 7 and first["skipped_lines"] == 1
    lines = con.execute("SELECT COUNT(*) FROM rebate_lines").fetchone()[0]
    ingest_zip(con, puf)
    assert con.execute("SELECT COUNT(*) FROM rebate_lines").fetchone()[0] == lines
    assert con.execute("SELECT COUNT(*) FROM ingest_runs").fetchone()[0] == 2


def test_reconcile_categorizes_every_case(con, puf):
    ingest_zip(con, puf)
    s = reconcile(con, 2024)
    got = {(r["filing_id"], r["check_name"]): r["category"] for r in con.execute("SELECT * FROM exceptions")}
    assert got == {
        (4, "adjusted_mlr"): "implausible_value",
        (5, "rebate"): "unexplained",
        (6, "rebate"): "nonpositive_premium",
    }
    assert s["merged_market_rebate_splits_verified"] == 1   # 30 + 50 = combined expected 80
    assert s["total_rebates_reported_usd"] == 45 + 30 + 50 + 99 + 3  # roll-up row (filing 7) excluded


def test_corrected_data_auto_resolves_and_regression_reopens(con, tmp_path, puf):
    ingest_zip(con, puf)
    reconcile(con, 2024)
    (tmp_path / "fixed").mkdir()
    fixed = build_zip(tmp_path / "fixed" / "mlr-public-use-file-2024.zip", default_data(mystery_rebate=10.0))
    ingest_zip(con, fixed)
    reconcile(con, 2024)
    e = con.execute("SELECT id, status FROM exceptions WHERE filing_id = 5").fetchone()
    assert e["status"] == "resolved"
    ingest_zip(con, puf)  # the bad figure comes back
    reconcile(con, 2024)
    assert con.execute("SELECT status FROM exceptions WHERE id = ?", (e["id"],)).fetchone()[0] == "open"
    trail = [r[0] for r in con.execute("SELECT to_status FROM exception_events WHERE exception_id = ? ORDER BY id",
                                       (e["id"],))]
    assert trail == ["open", "resolved", "open"]


def test_status_workflow_rejects_invalid_moves(con, puf):
    ingest_zip(con, puf)
    reconcile(con, 2024)
    eid = con.execute("SELECT id FROM exceptions WHERE filing_id = 5").fetchone()[0]
    set_status(con, eid, "investigating", "asked the filer")
    set_status(con, eid, "accepted", "confirmed late rebate adjustment")
    with pytest.raises(ValueError):
        set_status(con, eid, "investigating")
    with pytest.raises(KeyError):
        set_status(con, 12345, "resolved")


def test_peer_view_ranks_against_same_state_and_market(con, puf):
    ingest_zip(con, puf)
    rows = peer_view(con, 2024, "Clean Large Group")
    assert len(rows) == 1  # the Grand Total roll-up is not a state market
    r = rows[0]
    assert (r["state"], r["market"], r["peers"], r["rank_high_to_low"]) == ("IL", "Large group", 1, 1)
    assert r["peer_median_mlr"] == pytest.approx(0.71)


def test_workbook_uses_auditable_formulas(con, puf, tmp_path):
    ingest_zip(con, puf)
    s = reconcile(con, 2024)
    path = write_workbook(con, 2024, s, tmp_path / "report.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "By State", "Peer View", "Exceptions"]
    assert wb["By State"]["F2"].value.startswith("=IFERROR(D2/E2")
    ex = wb["Exceptions"]
    formulas = {ex.cell(r, 5).value: ex.cell(r, 16).value for r in range(2, ex.max_row + 1)}
    assert formulas["rebate"].startswith("=MAX(0,(N")
